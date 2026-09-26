"""Render a findings deck (.pptx) as one self-contained, interactive HTML page.

The page is a prebuilt React viewer (shadcn charts on Recharts, the house design system)
with the deck's content injected as JSON; building the viewer is the only step that needs node.

    /tmp/pptxenv/bin/python3 report/deck_html.py <deck.pptx> [...]   # writes <deck>.html beside each
    /tmp/pptxenv/bin/python3 report/deck_html.py --index <folder>     # every deck in it, plus index.html
    ... --lessons <talk.pptx>   # pin the talk deck's insights and hypothesis verdicts to their questions
    ... --out <folder>          # write the pages (and index.html) there instead of beside the decks

The deck is the source: native chart data, the side text, milestone lines and speaker notes
are read back out of the .pptx, so no chapter's deck.py changes to get a page. Slides that
share a tag prefix ("Q 2.01 · Answer", "Q 2.01 · By repo", "Q 2.01 · Change sets") become
tabs of one question, so a deck adds a view by adding a slide with the same prefix.

A chapter with a `chNN/book.md` reads as a book: its prose, with the charts it cites placed
inline as figures, then every question as an evidence appendix. The deck's own prose slides
are left out of that page, since the book replaces them.
"""
import argparse
import base64
import glob
import html
import json
import os
import re
import sys
from datetime import date, timedelta

from pptx import Presentation

import html_views
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU = 914400
C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
# The viewer is built from the private design system's components, so it lives in the gitignored
# .analysis/ beside the data rather than in this public repo.
VIEWER = os.environ.get("FINDINGS_VIEWER", os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(
    os.path.abspath(__file__)))), ".analysis", "viewer", "dist", "index.html"))
SIDE_X = 9.0
MONTHS = {m: i for i, m in enumerate(["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov",
                                      "Dec"], 1)}


def inch(v):
    return (v or 0) / EMU


def text_of(el):
    return "".join(t.text or "" for t in el.iter(A + "t"))


def paragraphs(shape):
    return [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]


# ------------------------------------------------------------------ charts

def _pts(cache, n=None):
    if cache is None:
        return []
    count = cache.find(C + "ptCount")
    n = int(count.get("val")) if count is not None else n or 0
    out = [None] * n
    for pt in cache.findall(C + "pt"):
        i = int(pt.get("idx"))
        v = pt.find(C + "v")
        if i >= len(out):
            out.extend([None] * (i + 1 - len(out)))
        out[i] = v.text if v is not None else None
    return out


def _color(sppr):
    if sppr is None:
        return None
    for tag in ("solidFill", "ln"):
        el = sppr.find(A + tag)
        if el is None:
            continue
        clr = el.find(".//" + A + "srgbClr")
        if clr is not None:
            return clr.get("val")
    return None


def _axis_title(ax):
    t = ax.find(C + "title") if ax is not None else None
    return text_of(t).strip() if t is not None else None


def series_name(tx):
    if tx is None:
        return ""
    v = tx.find(f".//{C}v")
    return (v.text if v is not None else text_of(tx)).strip()


def _weekly_labels(cats):
    """Timelines label every 4th week and blank the rest; put the week back on every one."""
    known = [(i, c) for i, c in enumerate(cats) if re.fullmatch(r"\d{1,2} [A-Z][a-z]{2}", c or "")]
    if not known or all((c or "").strip() for c in cats):
        return cats
    i0, c0 = known[0]
    d, m = c0.split()
    y = 2025 if MONTHS[m] == 12 and i0 == 0 else 2026
    base = date(y, MONTHS[m], int(d)) - timedelta(days=7 * i0)
    return [(base + timedelta(days=7 * i)).strftime("%-d %b") for i in range(len(cats))]


def read_chart(shape):
    cs = shape.chart._chartSpace
    plot_area = cs.find(f"{C}chart/{C}plotArea")
    ml = plot_area.find(f"{C}layout/{C}manualLayout")
    inner = {k: float(ml.find(C + k).get("val")) for k in ("x", "y", "w", "h")} if ml is not None else \
        {"x": 0.09, "y": 0.16, "w": 0.88, "h": 0.68}
    plot = next(el for el in plot_area if el.tag in (C + "barChart", C + "lineChart"))
    kind = "line" if plot.tag == C + "lineChart" else "bar"
    horizontal = kind == "bar" and plot.find(C + "barDir").get("val") == "bar"
    grouping = plot.find(C + "grouping")
    stacked = grouping is not None and grouping.get("val") in ("stacked", "percentStacked")
    series, cats = [], None
    for ser in plot.findall(C + "ser"):
        cat = ser.find(C + "cat")
        if cats is None and cat is not None:
            cache = cat.find(f".//{C}strCache")
            if cache is None:
                cache = cat.find(f".//{C}numCache")
            cats = [c or "" for c in _pts(cache)]
        vals = _pts(ser.find(f"{C}val//{C}numCache"), len(cats or []))
        points = {int(dp.find(C + "idx").get("val")): _color(dp.find(C + "spPr")) for dp in ser.findall(C + "dPt")}
        series.append({
            "name": series_name(ser.find(C + "tx")),
            "values": [None if v is None else float(v) for v in vals],
            "color": _color(ser.find(C + "spPr")),
            "pointColors": [points.get(i) for i in range(len(vals))] if points else None,
        })
    dl = plot.find(C + "dLbls")
    labels = dl is not None and dl.find(C + "showVal") is not None and dl.find(C + "showVal").get("val") == "1"
    val_ax, cat_ax = plot_area.find(C + "valAx"), plot_area.find(C + "catAx")
    fmt = val_ax.find(C + "numFmt").get("formatCode") if val_ax.find(C + "numFmt") is not None else "General"
    scaling = val_ax.find(C + "scaling")
    vmax = scaling.find(C + "max")
    vmin = scaling.find(C + "min")
    cats = _weekly_labels(cats or [])
    if horizontal:
        cats = cats[::-1]
        for s in series:
            s["values"] = s["values"][::-1]
            if s["pointColors"]:
                s["pointColors"] = s["pointColors"][::-1]
    return {
        "kind": kind, "horizontal": horizontal, "stacked": stacked, "cats": cats, "series": series,
        "labels": labels, "fmt": fmt, "legend": cs.find(f"{C}chart/{C}legend") is not None,
        "xTitle": _axis_title(cat_ax), "yTitle": _axis_title(val_ax),
        "max": float(vmax.get("val")) if vmax is not None else None,
        "min": float(vmin.get("val")) if vmin is not None else None,
        "frame": (inch(shape.left), inch(shape.top), inch(shape.width), inch(shape.height)), "inner": inner,
        "marks": [], "bands": [],
    }


def attach_overlays(charts, lines, rotated, rects):
    """Milestone lines and shaded spans are drawn shapes; turn each into a category position on its chart."""
    def owner(x):
        for ch in charts:
            fx, _, fw, _ = ch["frame"]
            if fx <= x <= fx + fw and not ch["horizontal"]:
                return ch
        return None

    def units(ch, x):
        fx, _, fw, _ = ch["frame"]
        left = fx + fw * ch["inner"]["x"]
        return (x - left) / (fw * ch["inner"]["w"]) * len(ch["cats"])

    used = set()
    for ln in lines:
        ch = owner(ln["x"])
        if not ch:
            continue
        near = min(((abs(t["cx"] - ln["x"]), i) for i, t in enumerate(rotated) if i not in used), default=None)
        name = ""
        if near and near[0] < 0.2:
            used.add(near[1])
            name = rotated[near[1]]["text"]
        ch["marks"].append({"at": round(units(ch, ln["x"]), 3), "name": name, "color": ln["color"]})
    for r in rects:
        ch = owner(r["x"] + r["w"] / 2)
        if ch:
            ch["bands"].append({"from": round(units(ch, r["x"]), 3), "to": round(units(ch, r["x"] + r["w"]), 3),
                                "name": r["text"], "color": r["color"]})


# ------------------------------------------------------------------ slides

def read_table(shape):
    return [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]


def read_slide(slide):
    s = {"layout": slide.slide_layout.name, "charts": [], "tables": [], "images": [], "texts": [], "notes": ""}
    lines, rotated, rects = [], [], []
    for sh in slide.shapes:
        x, y, w, h = inch(sh.left), inch(sh.top), inch(sh.width), inch(sh.height)
        if sh.has_chart:
            s["charts"].append(read_chart(sh))
        elif sh.shape_type == MSO_SHAPE_TYPE.TABLE:
            s["tables"].append(read_table(sh))
        elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            img = sh.image
            s["images"].append({"src": f"data:{img.content_type};base64,{base64.b64encode(img.blob).decode()}",
                                "alt": sh.name})
        elif sh.shape_type == MSO_SHAPE_TYPE.LINE:
            try:
                col = str(sh.line.color.rgb)
            except (AttributeError, TypeError):
                col = None
            lines.append({"x": x, "color": col})
        elif sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            try:
                col = str(sh.fill.fore_color.rgb)
            except (AttributeError, TypeError):
                col = None
            rects.append({"x": x, "w": w, "text": sh.text_frame.text.strip() if sh.has_text_frame else "",
                          "color": col})
        elif sh.has_text_frame and sh.text_frame.text.strip():
            if abs((sh.rotation or 0) - 270) < 1:
                rotated.append({"cx": x + w / 2, "text": sh.text_frame.text.strip()})
            else:
                s["texts"].append({"x": x, "y": y, "w": w, "paras": paragraphs(sh),
                                   "placeholder": sh.is_placeholder})
    attach_overlays(s["charts"], lines, rotated, rects)
    s["texts"].sort(key=lambda t: (t["y"], t["x"]))
    if slide.has_notes_slide:
        s["notes"] = slide.notes_slide.notes_text_frame.text.strip()
    return s


def classify(s):
    """Turn a slide's shapes into what it says: its kind, its tag and its parts."""
    t = s["texts"]
    if s["layout"] == "Title":
        return {"kind": "title", "title": t[0]["paras"][0] if t else "",
                "subtitle": " ".join(t[1]["paras"]) if len(t) > 1 else "", "notes": s["notes"]}
    if s["layout"] == "Section divider":
        return {"kind": "section", "kicker": " ".join(t[0]["paras"]) if t else "",
                "title": " ".join(t[1]["paras"]) if len(t) > 1 else "", "notes": s["notes"]}
    if s["charts"] or s["tables"] or s["images"]:
        side = [x for x in t if x["x"] >= SIDE_X]
        rest = [x for x in t if x["x"] < SIDE_X]
        tag = side[0]["paras"][0] if side else ""
        source = next((" ".join(x["paras"]) for x in side if x["paras"][0].upper().startswith("SOURCE")), "")
        body = [x for x in side[1:] if not x["paras"][0].upper().startswith("SOURCE")]
        return {"kind": "view", "tag": tag, "title": " ".join(body[0]["paras"]) if body else "",
                "points": [p for x in body[1:] for p in x["paras"]],
                "source": re.sub(r"^SOURCE\s*", "", source, flags=re.I),
                "extra": [p for x in rest for p in x["paras"]],
                "charts": s["charts"], "tables": s["tables"], "images": s["images"], "notes": s["notes"]}
    tag = t[0]["paras"][0] if t else ""
    if tag.upper().endswith("· QUESTION"):
        how, originally, cur = [], [], None
        for x in t[2:]:
            for p in x["paras"]:
                if p.upper() == "HOW WE ANSWER IT":
                    cur = how
                elif p.upper() == "ORIGINALLY ASKED":
                    cur = originally
                elif cur is not None:
                    cur.append(p)
        return {"kind": "question", "tag": tag, "question": " ".join(t[1]["paras"]) if len(t) > 1 else "",
                "how": " ".join(how), "originally": " ".join(originally), "notes": s["notes"]}
    return {"kind": "statement", "tag": tag, "headline": " ".join(t[1]["paras"]) if len(t) > 1 else "",
            "body": [p for x in t[2:] for p in x["paras"]], "notes": s["notes"]}


def split_tag(tag):
    parts = [p.strip() for p in tag.split("·")]
    if len(parts) == 1:
        return tag.strip(), ""
    return parts[0], " · ".join(parts[1:])


def group(slides):
    """Sections holding blocks; a block is one question (or topic) with its views as tabs."""
    head, sections, block = None, [], None
    for s in slides:
        k = s["kind"]
        if k == "title":
            head = s
            continue
        if k == "section" or not sections:
            sections.append({"kicker": s.get("kicker", ""), "title": s.get("title", ""),
                             "notes": s.get("notes", ""), "blocks": []})
            block = None
            if k == "section":
                continue
        blocks = sections[-1]["blocks"]
        if k == "question":
            key, _ = split_tag(s["tag"])
            block = {"key": key, "question": s, "views": [], "statements": []}
            blocks.append(block)
        elif k == "view":
            key, label = split_tag(s["tag"])
            if not block or block["key"].upper() != key.upper():
                block = {"key": key, "question": None, "views": [], "statements": []}
                blocks.append(block)
            s["label"] = label or "Answer"
            block["views"].append(s)
        else:
            blocks.append({"key": s["tag"], "question": None, "views": [], "statements": [s]})
            block = None
    return head or {"title": "", "subtitle": ""}, sections



# ------------------------------------------------------------------ lessons

Q_RE = re.compile(r"\bQ\s?(\d{1,2})\.(\d{2})\b")
VERDICT_RE = re.compile(r"^([A-Z][A-Z', ]+?)\s{2}·\s{2}(.*)$", re.S)


def qkey(v):
    m = Q_RE.search(v or "")
    return f"Q {int(m.group(1))}.{m.group(2)}" if m else None


def tone(verdict):
    v = verdict.upper()
    if v.startswith("HELD"):
        return "held"
    if v.startswith("PARTLY"):
        return "partly"
    if v.startswith("DID NOT"):
        return "not"
    return "open"


def read_lessons(path):
    """{"Q 9.03": {"insights": [...], "hypotheses": [...]}} from the evidence-edition talk deck."""
    out = {}
    def at(k):
        return out.setdefault(k, {"insights": [], "hypotheses": []})
    for slide in Presentation(path).slides:
        texts = [sh.text_frame for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        for i, tf in enumerate(texts):
            t = tf.text.strip()
            if t.startswith("HYPOTHESIS") and i + 1 < len(texts):
                paras = [p.text.strip() for p in texts[i + 1].paragraphs if p.text.strip()]
                m = VERDICT_RE.match(paras[-1]) if len(paras) > 1 else None
                if not m:
                    continue
                theme = t.split("\n", 1)[1].strip().title() if "\n" in t else ""
                evidence = re.split(r"\s+·\s+(?=Ch \d)", m.group(2).split("  ·  ")[0])[0].strip()
                refs = sorted({f"Q {int(a)}.{b}" for a, b in Q_RE.findall(m.group(2))})
                h = {"statement": " ".join(paras[:-1]), "verdict": m.group(1).strip(), "tone": tone(m.group(1)),
                     "evidence": evidence, "refs": refs, "theme": theme}
                for r in refs:
                    at(r)["hypotheses"].append(h)
            if t.upper().endswith("· FINDING"):
                key = qkey(t) or qkey(notes)
                body = next((x for x in texts if "THE INSIGHT" in x.text), None)
                if not key or body is None:
                    continue
                paras = [p.text.strip() for p in body.paragraphs if p.text.strip()]
                if "THE INSIGHT" not in paras:
                    continue
                insight = " ".join(paras[paras.index("THE INSIGHT") + 1:])
                if insight and insight not in at(key)["insights"]:
                    at(key)["insights"].append(insight)
    return out


def attach_lessons(sections, lessons):
    hits = 0
    for s in sections:
        for b in s["blocks"]:
            key = qkey(b["key"])
            if key and key in lessons:
                b["lessons"] = lessons[key]
                hits += 1
    return hits

# ------------------------------------------------------------------ html

def esc(v):
    return html.escape(v or "", quote=True)


def slim(sections):
    """Drop the slide geometry the page doesn't need once overlays are placed."""
    for s in sections:
        for b in s["blocks"]:
            for v in b["views"]:
                for ch in v["charts"]:
                    ch.pop("frame", None)
                    ch.pop("inner", None)
    return sections


def fill(template, title, payload):
    data = json.dumps(payload, separators=(",", ":"), ensure_ascii=False).replace("</", "<\\/")
    # str.replace, not format(): the bundled script is full of braces.
    return template.replace("__DECK_TITLE__", esc(title), 1).replace("__DECK_DATA__", data, 1)


def attach_views(sections, views, chapter):
    """Tabs a chapter's views.py saved: HTML-only views such as the other unit of work."""
    blocks = {qkey(b["key"]): b for s in sections for b in s["blocks"] if qkey(b["key"])}
    for key, extra in views.items():
        b = blocks.get(qkey(key))
        if b is None:
            print(f"{chapter}: no question {key} on the page; its {len(extra)} view(s) are dropped", file=sys.stderr)
            continue
        have = {v["label"].lower() for v in b["views"]}
        b["views"] += [v for v in extra if v["label"].lower() not in have]


# ------------------------------------------------------------------ book

BOOK_DIR = os.path.dirname(os.path.abspath(__file__))
FIG_RE = re.compile(r"^\[\[(.+?)\]\]$")
NUM_RE = re.compile(r"\$?\d[\d,]*(?:\.\d+)?%?")
# Numbers a reader can check without the data: question and chapter numbers, dates, years, URLs.
FREE_RE = re.compile(r"https?://\S+|\bQ\s?\d{1,2}\.\d{2}\b|\b(?:Chapters?|Ch|Part|Figure|Figures)\s+[\d.,– and]+"
                     r"|\b\d{1,2}\s+(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\b"
                     r"|\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\s+\d{1,2}\b|\b20\d\d\b")


def parse_book(text):
    """book.md: `## ` sections, `### ` subheads, `> ` callouts, `[[Q 9.04 · By repo | caption]]` figures."""
    sections, para = [{"title": "", "blocks": []}], []

    def flush():
        if para:
            sections[-1]["blocks"].append({"kind": "p", "text": " ".join(para)})
            para.clear()

    for line in re.sub(r"<!--.*?-->", "", text, flags=re.S).splitlines():
        line = line.strip()
        fig = FIG_RE.match(line)
        if line.startswith("# "):
            continue
        if line.startswith("## "):
            flush()
            sections.append({"title": line[3:].strip(), "blocks": []})
        elif not line:
            flush()
        elif line.startswith("### "):
            flush()
            sections[-1]["blocks"].append({"kind": "h3", "text": line[4:].strip()})
        elif line.startswith(">"):
            flush()
            quote = line.lstrip("> ").strip()
            blocks = sections[-1]["blocks"]
            if blocks and blocks[-1]["kind"] == "quote" and line != ">":
                blocks[-1]["text"] += " " + quote
            elif quote:
                blocks.append({"kind": "quote", "text": quote})
        elif fig:
            flush()
            ref, _, caption = fig.group(1).partition("|")
            sections[-1]["blocks"].append({"kind": "figure", "ref": ref.strip(), "caption": caption.strip()})
        else:
            para.append(line)
    flush()
    return sections


def find_view(sections, ref):
    """`Q 9.04` is its first view; `Q 9.04 · By repo` the view with that label; a diagram by its tag or label."""
    key, label = split_tag(ref)
    want = ref.upper()
    for s in sections:
        for b in s["blocks"]:
            for v in b["views"]:
                vlabel = (v.get("label") or "").upper()
                if qkey(ref):
                    hit = qkey(b["key"]) == qkey(key) and (not label or vlabel == label.upper())
                else:
                    hit = want in (vlabel, f"{b['key']} · {vlabel}".upper(), b["key"].upper())
                if hit:
                    return b, v
    return None, None


def bare(num):
    return num.strip("$%").replace(",", "")


def haystack(sections):
    parts = []
    for s in sections:
        for b in s["blocks"]:
            q = b.get("question") or {}
            parts += [q.get("question", ""), q.get("how", ""), q.get("notes", "")]
            for v in b["views"]:
                parts += [v["title"], v["source"], v["notes"], *v["points"], *v["extra"]]
                parts += [" ".join(r) for t in v["tables"] for r in t]
            for st in b["statements"]:
                parts += [st["headline"], st["notes"], *st["body"]]
    return {bare(n) for n in NUM_RE.findall(" ".join(parts))}


def build_book(book, sections, chapter):
    """Resolve each figure to its view and warn on anything the chapter's own data doesn't carry."""
    known, n, fig, out = haystack(sections), int(chapter[2:]), 0, []
    for s in book:
        blocks = []
        for blk in s["blocks"]:
            if blk["kind"] == "figure":
                b, v = find_view(sections, blk["ref"])
                if v is None:
                    print(f"{chapter}: book figure [[{blk['ref']}]] matches no view on the page", file=sys.stderr)
                    continue
                fig += 1
                blk.update({"number": f"{n}.{fig}", "key": b["key"], "view": v, "caption": blk["caption"] or v["title"]})
            t = blk.get("text") or blk.get("caption", "")
            for num in NUM_RE.findall(FREE_RE.sub(" ", t)):
                if bare(num) not in known and not (bare(num).isdigit() and int(bare(num)) <= 10):
                    print(f"{chapter}: book number {num!r} is not in the chapter's data: {t[:80]}…", file=sys.stderr)
            blocks.append(blk)
        if blocks or s["title"]:
            out.append({**s, "blocks": blocks})
    return out


def evidence(sections):
    """The appendix: every question, without the deck's prose and diagrams, which the book replaces."""
    out = []
    for s in sections:
        blocks = [b for b in s["blocks"] if qkey(b["key"])]
        if blocks:
            out.append({**s, "blocks": blocks})
    return out


def render(path, template, index_href=None, lessons=None, out_dir=None):
    prs = Presentation(path)
    head, sections = group([classify(read_slide(s)) for s in prs.slides])
    m = re.match(r"Ch (\d+)", os.path.basename(path))
    book = None
    if m:
        chapter = f"ch{int(m.group(1)):02d}"
        attach_views(sections, html_views.load(chapter), chapter)
        src = os.path.join(BOOK_DIR, chapter, "book.md")
        if os.path.exists(src):
            with open(src) as f:
                book = build_book(parse_book(f.read()), sections, chapter)
            sections = evidence(sections)
    if lessons:
        attach_lessons(sections, lessons)
    out = os.path.join(out_dir or os.path.dirname(path), os.path.splitext(os.path.basename(path))[0] + ".html")
    payload = {"head": head, "sections": slim(sections), "index": index_href, "book": book}
    with open(out, "w") as f:
        f.write(fill(template, head["title"], payload))
    return out, head


def render_index(folder, entries, template):
    chapters = [{"key": os.path.basename(h)[:5], "question": {"question": t["title"], "how": t["subtitle"],
                                                              "originally": "", "notes": ""},
                 "views": [], "statements": [], "href": os.path.basename(h)} for h, t in entries]
    payload = {"head": {"title": "Architecting a Software Factory", "subtitle": "One page per chapter: the argument first, then the evidence behind it."},
               "sections": [{"kicker": "Research findings", "title": "Chapters", "notes": "", "blocks": chapters}],
               "index": None}
    with open(os.path.join(folder, "index.html"), "w") as f:
        f.write(fill(template, "Research findings", payload))


def main(argv):
    ap = argparse.ArgumentParser(description="Render findings decks as HTML pages.")
    ap.add_argument("decks", nargs="*", help="deck .pptx files to render")
    ap.add_argument("--index", metavar="FOLDER", help="render every 'Ch *.pptx' in FOLDER, plus index.html")
    ap.add_argument("--lessons", metavar="TALK", help="talk deck whose insights and verdicts pin to questions")
    ap.add_argument("--out", metavar="FOLDER", help="write the pages there instead of beside the decks")
    args = ap.parse_args(argv)
    if not os.path.exists(VIEWER):
        sys.exit(f"no viewer build at {VIEWER}: run `npm run build` in its folder first")
    with open(VIEWER) as f:
        template = f.read()
    lessons = read_lessons(args.lessons) if args.lessons else None
    if args.index:
        entries = [render(p, template, "index.html", lessons, args.out)
                   for p in sorted(glob.glob(os.path.join(args.index, "Ch *.pptx")))]
        render_index(args.out or args.index, entries, template)
        for out, _ in entries:
            print(out)
        return
    for p in args.decks:
        print(render(p, template, lessons=lessons, out_dir=args.out)[0])


if __name__ == "__main__":
    main(sys.argv[1:])
